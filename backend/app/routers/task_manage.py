import json
import os
from concurrent.futures import ThreadPoolExecutor, TimeoutError as _FutureTimeout
from fastapi import APIRouter, Depends, HTTPException, Query, Header
from fastapi.responses import FileResponse
from sqlalchemy.orm import Session
from pydantic import BaseModel, Field
from typing import List, Optional
from datetime import datetime
from app.models.database import SessionLocal
from app.models.tables import (
    Task, Submission, Evaluation, LoginAccount, Teacher, Student, TaskClassRef,
)
from app.models.class_models import ClassMember, Class
from app.models.enterprise_models import JobPosition, Enterprise
from app.utils.ai_evaluator import (
    deepseek_client as _ds_client,
    is_ai_configured as _is_ai_configured,
    build_ai_misconfig_diagnosis as _build_ai_diag,
)
from app.utils.auth import decode_token

router = APIRouter(prefix="/api/tasks", tags=["任务管理"])

UPLOAD_DIR = "uploads/templates"
os.makedirs(UPLOAD_DIR, exist_ok=True)


def _normalize_skill_reqs(skill_reqs_raw):
    """解析并规范化岗位 skill_requirements。
    返回 [{name, weight_pct(0~100%), weight_raw, threshold, must, desc}]
    兼容两种 weight 单位：max<=1.5 视为归一化 0~1；否则视为百分比 0~100。
    """
    if not skill_reqs_raw:
        return []
    if isinstance(skill_reqs_raw, str):
        try:
            skill_reqs_raw = json.loads(skill_reqs_raw)
        except Exception:
            return []
    if not isinstance(skill_reqs_raw, list):
        return []
    parsed = []
    for s in skill_reqs_raw:
        if not isinstance(s, dict):
            continue
        nm = str(s.get("name", "")).strip()
        if not nm:
            continue
        w_raw = float(s.get("weight", 1.0) or 1.0)
        th = int(float(s.get("threshold", 60) or 60))
        must = bool(s.get("must", False))
        parsed.append({"name": nm, "weight_raw": w_raw, "threshold": th, "must": must})
    if not parsed:
        return []
    # 判断单位：所有 raw 都 <=1.5 → 归一化
    all_normalized = all(it["weight_raw"] <= 1.5 for it in parsed)
    if all_normalized:
        total = sum((it["weight_raw"] for it in parsed), 0.0) or 1.0
        for it in parsed:
            it["weight_pct"] = max(1, round(it["weight_raw"] / total * 100))
    else:
        # 百分比模式：先按和=100 归一（容忍教师/企业输入时合计略偏离）
        total = sum((it["weight_raw"] for it in parsed), 0.0) or 100.0
        for it in parsed:
            it["weight_pct"] = max(1, round(it["weight_raw"] / total * 100))
    # 修正合计，保证总和=100（差额塞到首个维度）
    delta = 100 - sum((it["weight_pct"] for it in parsed), 0)
    if delta != 0 and parsed:
        parsed[0]["weight_pct"] = max(1, parsed[0]["weight_pct"] + delta)
    # 生成描述：门槛 X 分 + 必选/加分
    for it in parsed:
        tag = "必达项" if it["must"] else "加分项"
        it["desc"] = f"岗位门槛 {it['threshold']}/100（{tag}）"
    return parsed


def _skill_reqs_to_task_criteria(skill_reqs_raw):
    """把岗位 skill_requirements 转化为 Task 保存所需的 3 组字段：
    (criteria_csv, criteria_weights_pct_csv, grading_dimensions_list)
    用于三方（企业/教师/AI）评分维度名称与权重完全对齐。
    """
    dims = _normalize_skill_reqs(skill_reqs_raw)
    if not dims:
        # 岗位没有技能门槛时回退一套默认技术岗维度，保证三方统一
        dims = [
            {"name": "代码质量", "weight_pct": 30, "desc": "岗位门槛 70/100（必达项）"},
            {"name": "功能完整性", "weight_pct": 30, "desc": "岗位门槛 75/100（必达项）"},
            {"name": "文档规范性", "weight_pct": 20, "desc": "岗位门槛 65/100（加分项）"},
            {"name": "岗位匹配度", "weight_pct": 20, "desc": "岗位门槛 70/100（必达项）"},
        ]
    criteria_names = ",".join([d["name"] for d in dims])
    criteria_weights = ",".join([str(d["weight_pct"]) for d in dims])
    grading_dimensions = [
        {"name": d["name"], "weight": d["weight_pct"], "desc": d["desc"]}
        for d in dims
    ]
    return criteria_names, criteria_weights, grading_dimensions


def get_db():
    db = SessionLocal()
    try:
        yield db
    finally:
        db.close()


def _account_to_teacher_id(db: Session, account_id: int):
    t = db.query(Teacher).filter(Teacher.account_id == int(account_id)).first()
    return t.id if t else None


def _account_to_student_id(db: Session, account_id: int):
    s = db.query(Student).filter(Student.account_id == int(account_id)).first()
    return s.id if s else None


def _student_id_to_account_id(db: Session, student_id: int):
    s = db.query(Student).filter(Student.id == int(student_id)).first()
    return s.account_id if s else None


def _require_role(roles):
    """返回 Depends 工厂：解码 token -> 查 login_accounts -> 按角色校验。
    兼容 teacher/student/enterprise（自动映射 mentor）
    """

    def _inner(
        token_query: Optional[str] = Query(None, alias="token"),
        authorization: Optional[str] = Header(None),
        db: Session = Depends(get_db),
    ):
        raw: Optional[str] = None
        if token_query:
            raw = token_query
        elif authorization and authorization.lower().startswith("bearer "):
            raw = authorization.split(None, 1)[1].strip()
        elif authorization:
            raw = authorization.strip()
        if not raw:
            raise HTTPException(401, "未登录或缺少 Token")
        payload = decode_token(raw)
        if not payload:
            raise HTTPException(401, "登录已过期或 Token 无效，请重新登录")
        user_id = payload.get("user_id") or payload.get("id") or payload.get("data", {}).get("user_id")
        if user_id is None:
            raise HTTPException(401, "登录已过期，请重新登录")
        acct = db.query(LoginAccount).filter(LoginAccount.id == int(user_id)).first()
        if not acct:
            raise HTTPException(401, "登录用户不存在，请重新登录")
        ext_role = "enterprise" if acct.role == "mentor" else (acct.role or "student")
        token_role = payload.get("role") or payload.get("data", {}).get("role")
        # 前端 token.role 可能是 enterprise；合并一下
        role = token_role or ext_role
        # 把 enterprise/mentor 统一成 enterprise 给外部比较
        check_role = "enterprise" if role in ("enterprise", "mentor") else role
        if roles:
            normalized_roles = set()
            for r in roles:
                if r in ("enterprise", "mentor"):
                    normalized_roles.update({"enterprise", "mentor"})
                else:
                    normalized_roles.add(r)
            if check_role not in normalized_roles and role not in normalized_roles:
                raise HTTPException(403, f"仅 {','.join(roles)} 可访问当前接口（当前角色 {role}）")
        return {"user_id": int(acct.id), "role": role, "account": acct, "payload": payload}

    return _inner


class GenerateRequest(BaseModel):
    title: str
    difficulty: str = "普通"


class AIFromJobRequest(BaseModel):
    job_id: int = Field(..., gt=0, description="企业岗位 ID")


class TaskCreate(BaseModel):
    title: str
    requirements: str
    criteria: str = "代码质量,功能完整性,文档规范性,界面设计"
    criteria_weights: str = "25,25,25,25"
    template_content: str = ""
    class_id: str = ""
    class_ids: Optional[List[int]] = None
    total_score: int = 100
    deadline: Optional[str] = None
    teacher_id: int = 0  # login_accounts.id；入库前映射为 teachers.id
    origin: str = "teacher_manual"
    linked_job_id: Optional[int] = None
    is_enterprise_project: int = 0
    ai_generated_job_title: str = ""


@router.post("/generate")
def generate_task(req: GenerateRequest):
    """AI 生成实训要求"""
    if not _is_ai_configured() or not _ds_client:
        diag = ""
        try:
            diag = _build_ai_diag()
        except Exception:
            diag = ""
        base_msg = (
            "AI 服务未配置：DEEPSEEK_API_KEY 为空或无效。\n\n"
            "⚠️  你是不是把 Key 填到了 .env.example？默认只读取 .env（.env.example 仅作模板）。\n"
            "正确做法（任选其一，都需要重启后端）：\n"
            "  方案A【推荐】: 复制 .env.example → .env，然后在 .env 里填 Key\n"
            "  方案B【本项目快捷】: 新版 config.py 已支持在没有 .env 时自动读 .env.example，只要 Ctrl+C 重启后端即可生效\n\n"
            "Key 申请地址：https://platform.deepseek.com/api_keys\n"
        )
        if diag:
            base_msg += "\n———————— 诊断报告（不含任何密钥）————————\n" + diag
        return {
            "success": False,
            "error": base_msg,
        }

    prompt = f"""你是软件实训课程设计专家。请根据以下信息生成实训任务：

任务标题：{req.title}
难度等级：{req.difficulty}
难度要求：{
    "3个评分维度，权重偏功能完整性，实训要求200字左右，模板4章节"
    if req.difficulty == "简单"
    else ("4个评分维度，权重均匀，实训要求250字左右，模板5章节"
          if req.difficulty == "普通"
          else "5-6个评分维度，权重偏代码质量和算法，实训要求300字左右，模板6章节")
}

请严格按JSON格式返回，必须包含 template 字段：
{{
    "requirements": "详细的实训要求，包含具体技术点、功能点、验收标准",
    "criteria": [
        {{"name": "维度名", "weight": 权重(百分比数字)}}
    ],
    "template": "成果物提交模板，每个章节用 ## 开头，如 ## 一、项目概述\\n请描述...\\n\\n## 二、功能实现\\n请列出..."
}}

只返回JSON，不要任何解释。"""

    executor = ThreadPoolExecutor(max_workers=1)

    def _call_ai():
        try:
            response = _ds_client.chat.completions.create(
                model="deepseek-chat",
                messages=[{"role": "user", "content": prompt}],
                temperature=0.7,
                timeout=15.0,
            )
            result = response.choices[0].message.content or ""
            if "```" in result:
                parts = result.split("```")
                result = parts[1] if len(parts) >= 2 else result
            if result.lstrip().startswith("json"):
                result = result.lstrip()[4:]
            result = result.strip()
            data = json.loads(result)
            if not isinstance(data.get("criteria"), list) or len(data["criteria"]) == 0:
                return False, "AI 返回结构异常：缺少 criteria 评分维度字段"
            if not isinstance(data.get("requirements"), str) or len(data["requirements"].strip()) < 8:
                return False, "AI 返回结构异常：缺少 requirements 实训要求字段"
            if not isinstance(data.get("template"), str) or len(data["template"].strip()) < 8:
                return False, "AI 返回结构异常：缺少 template 提交模板字段"
            return True, data
        except Exception as _e:
            msg = str(_e)
            if "401" in msg or "Unauthorized" in msg or "invalid_api_key" in msg.lower():
                msg = f"AI Key 无效(401)：{msg}. 请检查 .env 中的 DEEPSEEK_API_KEY 是否正确"
            elif "429" in msg or "rate limit" in msg.lower():
                msg = f"AI 限流(429)：{msg}. 请稍后再试或升级 API Key 额度"
            elif "500" in msg or "502" in msg or "503" in msg or "Bad Gateway" in msg:
                msg = f"AI 服务端错误({msg[:12]})：{msg}. 可稍后重试"
            elif "timeout" in msg.lower() or "timed out" in msg.lower():
                msg = f"AI 请求超时：{msg}"
            elif "connection" in msg.lower() or "network" in msg.lower() or "name or service not known" in msg.lower():
                msg = f"AI 网络连接失败：{msg}. 请检查服务器是否能访问 api.deepseek.com"
            return False, msg

    try:
        fut = executor.submit(_call_ai)
        ok, payload = fut.result(timeout=16)
        if ok:
            return {"success": True, "data": payload}
        return {"success": False, "error": str(payload)}
    except _FutureTimeout:
        return {
            "success": False,
            "error": (
                "AI 服务响应超时（15s）。可能原因：\n"
                "  1) 本机到 api.deepseek.com 的网络不通\n"
                "  2) DeepSeek 服务端临时慢\n"
                "请稍后重试。"
            ),
        }
    except Exception as e:
        return {"success": False, "error": f"调用失败：{str(e)}"}
    finally:
        executor.shutdown(wait=False, cancel_futures=True)


def _resolve_class_ids_from_req(req: TaskCreate) -> List[int]:
    """同时兼容前端传 class_ids: list[int] 或 class_id: 逗号分隔字符串。
    去重并转为 int 列表。"""
    result: set = set()
    if isinstance(req.class_ids, list):
        for x in req.class_ids:
            try:
                result.add(int(x))
            except Exception:
                pass
    if req.class_id and req.class_id.strip():
        for s in req.class_id.split(","):
            s = s.strip()
            if s.isdigit():
                result.add(int(s))
    return list(result)


@router.post("/")
def create_task(req: TaskCreate, db: Session = Depends(get_db)):
    deadline = None
    if req.deadline and req.deadline != "string" and req.deadline.strip():
        try:
            deadline = datetime.strptime(req.deadline, "%Y-%m-%d %H:%M")
        except Exception:
            try:
                deadline = datetime.strptime(req.deadline, "%Y-%m-%dT%H:%M:%S")
            except Exception:
                pass

    template_path = None
    if req.template_content and req.template_content.strip():
        from docx import Document
        doc = Document()
        doc.add_heading(f"实训成果物 - {req.title}", 0)
        for line in req.template_content.split('\n'):
            line = line.strip()
            if line.startswith('## ') or line.startswith('# '):
                doc.add_heading(line.replace('#', '').strip(), level=2)
            elif line.startswith('- ') or line.startswith('* '):
                doc.add_paragraph(line, style='List Bullet')
            elif line:
                doc.add_paragraph(line)

        filename = f"task_template_{datetime.now().strftime('%Y%m%d%H%M%S')}.docx"
        filepath = os.path.join(UPLOAD_DIR, filename)
        doc.save(filepath)
        template_path = filename

    # teacher_id (login_accounts.id) -> teachers.id
    created_by_teacher_id = None
    if req.teacher_id and req.teacher_id > 0:
        created_by_teacher_id = _account_to_teacher_id(db, req.teacher_id) or req.teacher_id

    origin_in = (req.origin or "").strip()
    if origin_in in ("enterprise_job_ai", "enterprise_job"):
        resolved_origin = "enterprise_job_ai"
    else:
        resolved_origin = "teacher_manual"

    class_ids = _resolve_class_ids_from_req(req)
    class_id_csv = ",".join(str(c) for c in class_ids) or None

    # ==============================================================
    # 三方评分维度统一（保存时强制执行）：
    #   当任务标记为「企业级项目实训」(is_enterprise_project=1) 且
    #   有关联岗位 (linked_job_id) 时，
    #   **忽略前端传来的 req.criteria / req.criteria_weights**，
    #   强制以岗位的 skill_requirements 作为评分维度唯一来源。
    #   这样即使教师端在 UI 上临时改了维度，保存时也会被纠正回岗位设置，
    #   保证企业评价维度名 ≡ 教师评价维度名 ≡ AI 评价维度名。
    # ==============================================================
    final_criteria = req.criteria
    final_criteria_weights = req.criteria_weights
    if req.is_enterprise_project and req.linked_job_id:
        try:
            jp = (
                db.query(JobPosition)
                .filter(JobPosition.id == int(req.linked_job_id))
                .first()
            )
            if jp is not None:
                c_csv, w_csv, _dims = _skill_reqs_to_task_criteria(jp.skill_requirements)
                if c_csv:
                    final_criteria = c_csv
                    final_criteria_weights = w_csv
        except Exception:
            # 任何异常都回退到前端值（宁可不一致，也不让保存失败）
            pass

    task = Task(
        title=req.title,
        requirements=req.requirements,
        criteria=final_criteria,
        criteria_weights=final_criteria_weights,
        template_path=template_path,
        class_id=class_id_csv,
        total_score=req.total_score,
        deadline=deadline,
        status="published",
        created_by=created_by_teacher_id if created_by_teacher_id else None,
        origin=resolved_origin,
        linked_job_id=req.linked_job_id,
        is_enterprise_project=1 if req.is_enterprise_project else 0,
        ai_generated_job_title=req.ai_generated_job_title or "",
    )
    db.add(task)
    db.flush()

    # 写 task_class_ref
    for cid in class_ids:
        if cid and cid > 0:
            ref = TaskClassRef(task_id=task.id, class_id=int(cid))
            db.add(ref)

    db.commit()
    db.refresh(task)
    return {"success": True, "data": {"id": task.id, "title": task.title,
                                      "is_enterprise_project": bool(task.is_enterprise_project),
                                      "origin": task.origin,
                                      "class_ids": class_ids}}


@router.get("/enterprise-jobs")
def list_enterprise_jobs_for_teachers(
    keyword: str = "",
    status: str = "open",
    page: int = 1,
    page_size: int = 50,
    _user: dict = Depends(_require_role(["teacher", "admin", "enterprise"])),
    db: Session = Depends(get_db),
):
    q = db.query(JobPosition, Enterprise).outerjoin(Enterprise, Enterprise.id == JobPosition.enterprise_id)
    if status in ("open", "draft", "closed"):
        q = q.filter(JobPosition.status == status)
    else:
        q = q.filter(JobPosition.status == "open")
    if keyword:
        kw = f"%{keyword.strip()}%"
        q = q.filter(
            (JobPosition.title.like(kw))
            | (JobPosition.city.like(kw))
            | (JobPosition.tags.like(kw))
            | (JobPosition.requirements.like(kw))
            | (Enterprise.name.like(kw))
            | (Enterprise.short_name.like(kw))
        )
    total = q.count()
    rows = (
        q.order_by(JobPosition.updated_at.desc())
        .offset((page - 1) * page_size)
        .limit(page_size)
        .all()
    )
    items = []
    for jp, ent in rows:
        tag_list = [t for t in (jp.tags or "").split(",") if t]
        # 把岗位技能门槛也返回给教师端：
        #   - 用于前端展示"评分维度已与岗位技能门槛对齐"的提示横幅
        #   - 教师"跳过AI直接编辑"时可直接拿来初始化评分维度
        aligned_dims = _normalize_skill_reqs(jp.skill_requirements)
        items.append({
            "id": jp.id,
            "enterprise_id": jp.enterprise_id,
            "enterprise_name": (ent.short_name or ent.name) if ent else "企业",
            "enterprise_logo": ent.logo if ent else "",
            "title": jp.title,
            "level": jp.level or "初级",
            "salary_range": jp.salary_range or "",
            "city": jp.city or "",
            "summary": (jp.description or "")[:120],
            "tag_list": tag_list,
            "status": jp.status,
            "created_at": str(jp.created_at) if jp.created_at else "",
            # 供教师端对齐评分维度
            "skill_requirements": aligned_dims,
            "job_type": getattr(jp, "job_type", None) or "技术岗",
            "criteria_aligned": bool(aligned_dims),
        })
    return {"success": True, "total": total, "page": page, "page_size": page_size, "list": items}


@router.post("/ai-from-job")
def ai_generate_task_from_job(
    req: AIFromJobRequest,
    _u: dict = Depends(_require_role(["teacher", "admin"])),
    db: Session = Depends(get_db),
):
    if not _is_ai_configured() or not _ds_client:
        diag = ""
        try:
            diag = _build_ai_diag()
        except Exception:
            diag = ""
        base_msg = (
            "AI 服务未配置：DEEPSEEK_API_KEY 为空或无效。请先在后端 .env 中配置有效的 DeepSeek API Key 并重启后端。\n"
            "Key 申请：https://platform.deepseek.com/api_keys"
        )
        if diag:
            base_msg += "\n\n诊断报告：\n" + diag
        raise HTTPException(status_code=502, detail=base_msg)

    row = (
        db.query(JobPosition, Enterprise)
        .outerjoin(Enterprise, Enterprise.id == JobPosition.enterprise_id)
        .filter(JobPosition.id == req.job_id)
        .first()
    )
    if not row:
        raise HTTPException(404, "所选企业岗位不存在或已被删除")
    jp, ent = row
    if jp.status != "open":
        raise HTTPException(400, "该岗位未处于招聘中，请选择其它岗位")

    skill_reqs = jp.skill_requirements or []
    if isinstance(skill_reqs, str):
        try:
            skill_reqs = json.loads(skill_reqs)
        except Exception:
            skill_reqs = []
    if not isinstance(skill_reqs, list):
        skill_reqs = []
    skill_dims_str = "\n".join([
        f"  - 维度「{s.get('name','')}」：门槛 {s.get('threshold',60)} 分，权重 {s.get('weight','?')}，{'必须' if s.get('must') else '加分项'}"
        for s in skill_reqs if isinstance(s, dict)
    ]) or "  - （未设置）"
    tag_str = "、".join([t for t in (jp.tags or "").split(",") if t]) or "（未设置）"

    user_prompt = f"""你是资深软件工程实训课程设计师，请基于一份企业真实招聘岗位，设计一份适合职业院校/本科软件实训的「企业级项目实训任务」。
注意：这是学生需要提交代码、项目文档的真实实训，不是理论题目。

企业信息：
  - 企业：{(ent.short_name or ent.name) if ent else '企业'}
  - 岗位名称：{jp.title}
  - 级别：{jp.level or '初级'}
  - 城市：{jp.city or '不限'}
  - 薪资范围：{jp.salary_range or '面议'}
  - 岗位技能标签：{tag_str}
岗位原文：
  - 描述：{jp.description or '略'}
  - 岗位职责：{jp.responsibilities or '略'}
  - 任职要求：{jp.requirements or '略'}
  - 岗位技能门槛维度：
{skill_dims_str}

请严格 JSON 返回，只返回 JSON，不要任何解释：
{{
  "title": "实训任务标题（要具象，含项目名，如「{jp.title} · XX 项目实战」）",
  "description": "实训任务简介，150~220 字，讲清楚这是一个什么企业级项目，业务背景是什么，需要学生交付什么",
  "requirements": [
    "功能要求第 1 条（具体）",
    "功能要求第 2 条",
    "... 共 4~8 条，要具体，和该岗位的职责对应"
  ],
  "deliverables": [
    "交付物第 1 项：项目源代码（需 push 至 Gitee/GitHub 仓库或打包 .zip）",
    "交付物第 2 项：项目设计与说明文档（格式要求：需求分析、架构图、数据库设计、关键实现）",
    "交付物第 3 项：README.md（运行说明、启动步骤、示例账号）",
    "交付物第 4 项：演示视频或截图（可选）"
  ],
  "grading_dimensions": [
    {{"name":"代码质量","weight":25,"desc":"规范、模块化、异常处理"}},
    {{"name":"功能完整性","weight":30,"desc":"覆盖关键功能点"}},
    {{"name":"项目文档","weight":20,"desc":"结构清晰、格式规范"}},
    {{"name":"UI/交互","weight":10,"desc":"界面美观、交互流畅"}},
    {{"name":"岗位匹配度","weight":15,"desc":"实现技能与岗位门槛维度对应"}}
  ]
}}
约束：
1) grading_dimensions 所有权重之和必须 =100
2) requirements 必须和该岗位的职责/任职要求强相关，不能生成通用 CRUD
3) description 必须具体，禁止空泛模板
4) deliverables 必须明确"文件/格式"，因为学生要提交到平台做代码 + 文档审核
"""

    def _call_ai():
        try:
            resp = _ds_client.chat.completions.create(
                model="deepseek-chat",
                messages=[{"role": "user", "content": user_prompt}],
                temperature=0.78,
                response_format={"type": "json_object"},
                timeout=35.0,
            )
            raw = resp.choices[0].message.content or ""
            if "```" in raw:
                parts = raw.split("```")
                raw = parts[1] if len(parts) >= 2 else raw
                if raw.lstrip().startswith("json"):
                    raw = raw.lstrip()[4:]
            data = json.loads(raw.strip())
            for k in ("title", "description", "requirements", "deliverables", "grading_dimensions"):
                if k not in data:
                    return False, f"AI 返回结构异常：缺少字段 {k}"
            if not isinstance(data["requirements"], list) or len(data["requirements"]) < 2:
                return False, "AI 返回结构异常：requirements 至少 2 条"
            if not isinstance(data["deliverables"], list) or len(data["deliverables"]) < 2:
                return False, "AI 返回结构异常：deliverables 至少 2 条"
            if not isinstance(data["grading_dimensions"], list) or len(data["grading_dimensions"]) < 2:
                return False, "AI 返回结构异常：grading_dimensions 至少 2 条"
            return True, data
        except Exception as _e:
            msg = str(_e)
            if "401" in msg or "Unauthorized" in msg or "invalid_api_key" in msg.lower():
                msg = f"AI Key 无效(401)：{msg}。请检查 DEEPSEEK_API_KEY 是否正确"
            elif "429" in msg or "rate limit" in msg.lower():
                msg = f"AI 限流(429)：{msg}。请稍后再试或升级额度"
            elif "500" in msg or "502" in msg or "503" in msg:
                msg = f"AI 服务端错误：{msg}"
            elif "timeout" in msg.lower() or "timed out" in msg.lower():
                msg = f"AI 请求超时(35s)：{msg}"
            elif "connection" in msg.lower() or "network" in msg.lower():
                msg = f"AI 网络连接失败：{msg}。请检查服务器能否访问 api.deepseek.com"
            return False, msg

    executor = ThreadPoolExecutor(max_workers=1)
    try:
        fut = executor.submit(_call_ai)
        ok, payload = fut.result(timeout=37)
        if not ok:
            raise HTTPException(status_code=502, detail=str(payload))
        dims = []
        seen_names: set = set()
        for g in (payload.get("grading_dimensions") or []):
            if not isinstance(g, dict):
                continue
            nm = str(g.get("name", "")).strip()
            if not nm or nm in seen_names:
                continue
            seen_names.add(nm)
            w = int(float(g.get("weight", 25)))
            if w < 0:
                w = 0
            desc = str(g.get("desc", "")).strip()
            dims.append({"name": nm, "weight": w, "desc": desc})
        total_w = sum((d["weight"] for d in dims), 0) or 1
        for d in dims:
            d["weight"] = round(d["weight"] * 100 / total_w)
        if dims:
            delta = 100 - sum((d["weight"] for d in dims), 0)
            dims[-1]["weight"] += delta
            if dims[-1]["weight"] < 0:
                dims[-1]["weight"] = 0

        requirements_joined = "；\n".join([str(r).strip() for r in (payload.get("requirements") or []) if str(r).strip()])
        deliverables_joined = "\n".join([f"## {i+1}. {str(d).strip()}" for i, d in enumerate(payload.get("deliverables") or []) if str(d).strip()])

        # ==============================================================
        # 三方评分维度统一：**无视 AI 生成的 grading_dimensions**，
        # 强制使用岗位 skill_requirements 作为唯一准绳，保证：
        #   企业端评价维度名 == 教师端评价维度名 == AI 评价维度名
        #   教师端 criteria/criteria_weights == 岗位技能门槛
        # ==============================================================
        job_criteria_csv, job_weights_csv, job_dims = _skill_reqs_to_task_criteria(jp.skill_requirements)
        dims = job_dims
        criteria_names = job_criteria_csv
        criteria_weights = job_weights_csv

        template_content = f"## 企业级项目实训 - 成果物提交模板\n\n### 项目：{payload.get('title','')}\n企业岗位：{jp.title}\n\n## 1. 项目概述\n请用 300 字描述项目背景、目标与核心功能模块。\n\n{deliverables_joined}\n\n## 附：评分维度说明（与企业岗位技能门槛一致）\n" + "\n".join([f"- {d['name']}（{d['weight']}分）：{d['desc']}" for d in dims])
        full_requirements = f"### 实训简介\n{payload.get('description','')}\n\n### 功能要求\n" + "\n".join([f"{i+1}. {str(r).strip()}" for i, r in enumerate(payload.get("requirements") or []) if str(r).strip()])

        # 把岗位原始 skill_requirements 也回传给前端，用于展示"维度已对齐"提示
        aligned_dims_for_fe = _normalize_skill_reqs(jp.skill_requirements) or [
            {"name": d["name"], "threshold": 70, "must": False,
             "weight_pct": d["weight"], "desc": d["desc"]}
            for d in dims
        ]

        return {
            "success": True,
            "job": {
                "id": jp.id,
                "enterprise_id": jp.enterprise_id,
                "enterprise_name": (ent.short_name or ent.name) if ent else "企业",
                "title": jp.title,
                "level": jp.level or "",
                "city": jp.city or "",
                "salary_range": jp.salary_range or "",
                "tag_list": [t for t in (jp.tags or "").split(",") if t],
                # 供前端展示"维度已与岗位技能门槛对齐"
                "skill_requirements": aligned_dims_for_fe,
                "criteria_aligned": True,
            },
            "draft": {
                "title": payload.get("title", ""),
                "description": payload.get("description", ""),
                "requirements_text": full_requirements,
                "requirements_list": [str(r).strip() for r in (payload.get("requirements") or []) if str(r).strip()],
                "deliverables_list": [str(d).strip() for d in (payload.get("deliverables") or []) if str(d).strip()],
                "grading_dimensions": dims,
                "criteria": criteria_names,
                "criteria_weights": criteria_weights,
                "template_content": template_content,
                "criteria_aligned_to_job": True,
            },
        }
    except _FutureTimeout:
        raise HTTPException(status_code=504, detail="AI 服务响应超时(35s)。请稍后重试，或检查后端到 api.deepseek.com 网络")
    finally:
        executor.shutdown(wait=False, cancel_futures=True)


@router.get("/{task_id}/template")
def download_template(task_id: int, db: Session = Depends(get_db)):
    task = db.query(Task).filter(Task.id == task_id).first()
    if not task or not task.template_path:
        raise HTTPException(404, "该任务没有模板")
    filepath = os.path.join(UPLOAD_DIR, task.template_path)
    if not os.path.exists(filepath):
        raise HTTPException(404, "模板文件不存在")
    return FileResponse(
        filepath,
        filename=f"{task.title}_提交模板.docx",
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )


def _get_task_class_ids(db: Session, task_ids: List[int]) -> dict:
    """批量查 tasks <-> class_ids（优先用 task_class_ref，回退 class_id 列）。返回 {task_id: [class_id,...]}"""
    res: dict = {int(t): [] for t in task_ids}
    if not task_ids:
        return res
    refs = (
        db.query(TaskClassRef.task_id, TaskClassRef.class_id)
        .filter(TaskClassRef.task_id.in_(list(set(task_ids))))
        .all()
    )
    for tid, cid in refs:
        res.setdefault(int(tid), []).append(int(cid))
    # 空的从 class_id 列补（兼容老数据）
    for tid in list(res.keys()):
        if res[tid]:
            continue
        t = db.query(Task).filter(Task.id == tid).first()
        if t and t.class_id:
            for s in str(t.class_id).split(","):
                s = s.strip()
                if s.isdigit():
                    res[tid].append(int(s))
    return res


@router.get("/my")
def get_my_tasks(teacher_id: int = 0, db: Session = Depends(get_db)):
    # teacher_id 约定是 login_accounts.id -> 映射为 teachers.id
    teacher_pk = _account_to_teacher_id(db, teacher_id) if teacher_id else None
    if teacher_pk:
        q_filter = Task.created_by == teacher_pk
    else:
        # 没映射到也允许按原 id 查（兼容老数据）
        q_filter = Task.created_by == teacher_id
    tasks = db.query(Task).filter(q_filter).order_by(Task.created_at.desc()).all()

    task_class_ids_map = _get_task_class_ids(db, [t.id for t in tasks])
    all_class_ids = set()
    for lst in task_class_ids_map.values():
        for c in lst:
            all_class_ids.add(int(c))
    job_ids = set()
    for t in tasks:
        if getattr(t, "linked_job_id", None):
            job_ids.add(int(t.linked_job_id))

    class_map: dict = {}
    if all_class_ids:
        classes = db.query(Class).filter(Class.id.in_(list(all_class_ids))).all()
        class_map = {c.id: c.name for c in classes}
    job_map: dict = {}
    if job_ids:
        for j in (
            db.query(JobPosition, Enterprise)
            .outerjoin(Enterprise, Enterprise.id == JobPosition.enterprise_id)
            .filter(JobPosition.id.in_(list(job_ids)))
            .all()
        ):
            jp, ent = j
            job_map[int(jp.id)] = {
                "id": jp.id,
                "title": jp.title,
                "enterprise_name": (ent.short_name or ent.name) if ent else "企业",
            }

    data_list = []
    for t in tasks:
        ent_label = None
        if getattr(t, "is_enterprise_project", 0):
            j = job_map.get(int(getattr(t, "linked_job_id", 0) or 0))
            job_title_snapshot = getattr(t, "ai_generated_job_title", "") or ""
            if j:
                ent_label = {"job_title": j["title"], "enterprise_name": j["enterprise_name"]}
            elif job_title_snapshot:
                ent_label = {"job_title": job_title_snapshot, "enterprise_name": "企业"}
        cids = task_class_ids_map.get(t.id, [])
        data_list.append({
            "id": t.id,
            "title": t.title,
            "requirements": t.requirements,
            "criteria": t.criteria,
            "criteria_weights": t.criteria_weights or "",
            "template_path": t.template_path,
            "class_ids": cids,
            "class_names": ", ".join([class_map.get(c, str(c)) for c in cids]) if cids else "不限",
            "total_score": t.total_score,
            "status": t.status,
            "deadline": str(t.deadline) if t.deadline else "",
            "created_at": str(t.created_at),
            "origin": getattr(t, "origin", "teacher_manual") or "teacher_manual",
            "linked_job_id": getattr(t, "linked_job_id", None),
            "is_enterprise_project": bool(getattr(t, "is_enterprise_project", 0)),
            "ai_generated_job_title": getattr(t, "ai_generated_job_title", "") or "",
            "enterprise_label": ent_label,
        })
    return {"success": True, "data": data_list}


@router.get("/pending")
def get_pending_tasks(student_id: int = 0, db: Session = Depends(get_db)):
    # student_id 约定是 login_accounts.id -> students.id
    stu_pk = _account_to_student_id(db, student_id) if student_id else None
    if stu_pk:
        student_fk = stu_pk
    else:
        student_fk = student_id

    members = db.query(ClassMember).filter(ClassMember.student_id == student_fk).all()
    class_ids = [m.class_id for m in members]
    if not class_ids:
        return {"success": True, "data": []}

    all_tasks = db.query(Task).filter(Task.status == "published").order_by(Task.created_at.desc()).all()
    task_class_ids_map = _get_task_class_ids(db, [t.id for t in all_tasks])

    job_ids = {int(t.linked_job_id) for t in all_tasks if getattr(t, "linked_job_id", None)}
    job_map: dict = {}
    if job_ids:
        for j in (
            db.query(JobPosition, Enterprise)
            .outerjoin(Enterprise, Enterprise.id == JobPosition.enterprise_id)
            .filter(JobPosition.id.in_(list(job_ids)))
            .all()
        ):
            jp, ent = j
            job_map[int(jp.id)] = {
                "id": jp.id,
                "title": jp.title,
                "enterprise_name": (ent.short_name or ent.name) if ent else "企业",
            }

    result = []
    for t in all_tasks:
        task_class_ids = task_class_ids_map.get(t.id, [])
        # 任务没有绑定班级视为"不限"，对所有学生可见
        if task_class_ids and not any(c in class_ids for c in task_class_ids):
            continue

        submitted = db.query(Submission).filter(
            Submission.task_id == t.id,
            Submission.student_id == student_fk
        ).first()

        ent_label = None
        if getattr(t, "is_enterprise_project", 0):
            j = job_map.get(int(getattr(t, "linked_job_id", 0) or 0))
            job_title_snapshot = getattr(t, "ai_generated_job_title", "") or ""
            if j:
                ent_label = {"job_title": j["title"], "enterprise_name": j["enterprise_name"]}
            elif job_title_snapshot:
                ent_label = {"job_title": job_title_snapshot, "enterprise_name": "企业"}

        result.append({
            "id": t.id,
            "title": t.title,
            "requirements": t.requirements,
            "criteria": t.criteria,
            "criteria_weights": t.criteria_weights or "",
            "template_path": t.template_path,
            "class_id": ",".join(str(c) for c in task_class_ids),
            "class_ids": task_class_ids,
            "total_score": t.total_score,
            "deadline": str(t.deadline) if t.deadline else "",
            "submitted": submitted is not None,
            "submission_id": submitted.id if submitted else None,
            "created_at": str(t.created_at) if t.created_at else "",
            "origin": getattr(t, "origin", "teacher_manual") or "teacher_manual",
            "linked_job_id": getattr(t, "linked_job_id", None),
            "is_enterprise_project": bool(getattr(t, "is_enterprise_project", 0)),
            "ai_generated_job_title": getattr(t, "ai_generated_job_title", "") or "",
            "enterprise_label": ent_label,
        })

    return {"success": True, "data": result}


@router.get("/{task_id}/detail")
def task_detail(task_id: int, db: Session = Depends(get_db)):
    task = db.query(Task).filter(Task.id == task_id).first()
    if not task:
        raise HTTPException(404, "任务不存在")

    task_class_ids_map = _get_task_class_ids(db, [task.id])
    task_class_ids = task_class_ids_map.get(task.id, [])

    raw_subs = (
        db.query(Submission)
        .filter(Submission.task_id == task_id)
        .order_by(Submission.created_at.desc(), Submission.id.desc())
        .all()
    )
    submissions = []
    seen_student = set()
    for s in raw_subs:
        key = s.student_id
        if key is None:
            key = f"__anon_{s.id}"
        if key in seen_student:
            continue
        seen_student.add(key)
        submissions.append(s)

    class_map = {}
    all_class_ids_for_subs = set()
    for sub in submissions:
        if sub.class_id:
            for cid in str(sub.class_id).split(','):
                cid = cid.strip()
                if cid.isdigit():
                    all_class_ids_for_subs.add(int(cid))
    if all_class_ids_for_subs:
        classes = db.query(Class).filter(Class.id.in_(list(all_class_ids_for_subs))).all()
        class_map = {c.id: c.name for c in classes}
    # 任务关联的班级也加进来
    if task_class_ids:
        if not class_map:
            classes = db.query(Class).filter(Class.id.in_(task_class_ids)).all()
            class_map = {c.id: c.name for c in classes}
        else:
            missing = [c for c in task_class_ids if c not in class_map]
            if missing:
                extra = db.query(Class).filter(Class.id.in_(missing)).all()
                for c in extra:
                    class_map[c.id] = c.name

    ent_label = None
    linked_job_id = getattr(task, "linked_job_id", None)
    if getattr(task, "is_enterprise_project", 0) and linked_job_id:
        row = (
            db.query(JobPosition, Enterprise)
            .outerjoin(Enterprise, Enterprise.id == JobPosition.enterprise_id)
            .filter(JobPosition.id == int(linked_job_id))
            .first()
        )
        if row:
            jp, ent = row
            ent_label = {
                "job_title": jp.title,
                "enterprise_name": (ent.short_name or ent.name) if ent else "企业",
            }
    if not ent_label and getattr(task, "is_enterprise_project", 0):
        title_snapshot = getattr(task, "ai_generated_job_title", "") or ""
        if title_snapshot:
            ent_label = {"job_title": title_snapshot, "enterprise_name": "企业"}

    student_list = []
    for sub in submissions:
        ai_eval = db.query(Evaluation).filter(
            Evaluation.submission_id == sub.id,
            Evaluation.evaluator_type == "ai"
        ).first()
        teacher_eval = db.query(Evaluation).filter(
            Evaluation.submission_id == sub.id,
            Evaluation.evaluator_type == "teacher"
        ).first()

        # sub.student_id -> students.id，要拿 real_name / user_number 走 students.account_id -> login_accounts
        real_name = "未知"
        user_number = ""
        if sub.student_id:
            stu = db.query(Student).filter(Student.id == int(sub.student_id)).first()
            if stu:
                real_name = stu.real_name or "未知"
                user_number = stu.student_no or ""
                if real_name in ("", "未知"):
                    acct = db.query(LoginAccount).filter(LoginAccount.id == stu.account_id).first()
                    if acct:
                        real_name = acct.username or "未知"

        class_names = "不限"
        if sub.student_id:
            memberships = db.query(ClassMember).filter(ClassMember.student_id == sub.student_id).all()
            if memberships:
                cnames = [class_map.get(m.class_id, str(m.class_id)) for m in memberships if m.class_id in class_map]
                if cnames:
                    class_names = "、".join(cnames)

        student_list.append({
            "submission_id": sub.id,
            "student_id": sub.student_id,
            "account_id": _student_id_to_account_id(db, sub.student_id) if sub.student_id else None,
            "student_name": real_name,
            "student_number": user_number,
            "filename": sub.filename,
            "submitted_at": str(sub.created_at) if sub.created_at else "",
            "class_name": class_names,
            "ai_score": float(ai_eval.total_score) if ai_eval else None,
            "ai_scores": ai_eval.dimension_scores if ai_eval else None,
            "ai_comment": ai_eval.comment if ai_eval else None,
            "ai_steps": ai_eval.step_completeness if ai_eval else None,
            "ai_issues": ai_eval.logic_issues if ai_eval else None,
            "teacher_score": float(teacher_eval.total_score) if teacher_eval else None,
            "is_scored": teacher_eval is not None,
        })

    unscored_count = sum(1 for s in student_list if not s["is_scored"])
    next_unscored_id = None
    for s in student_list:
        if not s["is_scored"]:
            next_unscored_id = s["submission_id"]
            break

    return {
        "success": True,
        "data": {
            "task": {
                "id": task.id,
                "title": task.title,
                "requirements": task.requirements,
                "criteria": task.criteria,
                "criteria_weights": task.criteria_weights or "",
                "template_path": task.template_path,
                "class_ids": task_class_ids,
                "class_names": ", ".join([class_map.get(c, str(c)) for c in task_class_ids]) if task_class_ids else "不限",
                "origin": getattr(task, "origin", "teacher_manual") or "teacher_manual",
                "linked_job_id": getattr(task, "linked_job_id", None),
                "is_enterprise_project": bool(getattr(task, "is_enterprise_project", 0)),
                "ai_generated_job_title": getattr(task, "ai_generated_job_title", "") or "",
                "enterprise_label": ent_label,
            },
            "submissions": student_list,
            "unscored_count": unscored_count,
            "next_unscored_id": next_unscored_id
        }
    }


@router.delete("/{task_id}")
def delete_task(task_id: int, db: Session = Depends(get_db)):
    task = db.query(Task).filter(Task.id == task_id).first()
    if not task:
        raise HTTPException(404, "任务不存在")
    db.query(TaskClassRef).filter(TaskClassRef.task_id == task_id).delete()
    submissions = db.query(Submission).filter(Submission.task_id == task_id).all()
    for sub in submissions:
        db.query(Evaluation).filter(Evaluation.submission_id == sub.id).delete()
        try:
            from app.models.enterprise_models import EnterpriseEvaluation
            db.query(EnterpriseEvaluation).filter(EnterpriseEvaluation.submission_id == sub.id).delete()
        except Exception:
            pass
        db.delete(sub)
    db.delete(task)
    db.commit()
    return {"success": True, "message": "任务已删除"}


@router.get("/download/{submission_id}")
def download_submission(submission_id: int, db: Session = Depends(get_db)):
    submission = db.query(Submission).filter(Submission.id == submission_id).first()
    if not submission:
        raise HTTPException(404, "提交记录不存在")
    if not submission.file_path or not os.path.exists(submission.file_path):
        raise HTTPException(404, "文件不存在")
    return FileResponse(
        submission.file_path,
        filename=submission.filename,
        media_type="application/octet-stream"
    )
